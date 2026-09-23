# -*- coding: utf-8 -*-
"""
Apresentação (.pptx) da Malha Viária
====================================

Gera uma apresentação 16:9 no padrão visual da Eixo (fundo grafite, destaque
amarelo, Arial) a partir dos quantitativos já calculados pelo relatório
(`rel` de relatorio_malha_viaria.gerar_relatorio).

Os gráficos são NATIVOS do PowerPoint (editáveis: dá para clicar e mudar
cores/rótulos), não imagens.
"""

import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from relatorio_malha_viaria import (
    filtrar_bairros_validos, formatar_inteiro_br, formatar_numero_br,
    formatar_pct, get_cor_pavimentacao,
)

# --------------------------------------------------------------------------
# Identidade visual
# --------------------------------------------------------------------------
AMARELO = RGBColor(0xF5, 0xC8, 0x00)
GRAFITE = RGBColor(0x2B, 0x2B, 0x2B)
GRAFITE_MEDIO = RGBColor(0x3D, 0x3D, 0x3D)
CARD_ESCURO = RGBColor(0x45, 0x45, 0x45)
CINZA_CLARO = RGBColor(0xF2, 0xF2, 0xF2)
CINZA_TEXTO = RGBColor(0x6B, 0x6B, 0x6B)
CINZA_SUAVE = RGBColor(0xCC, 0xCC, 0xCC)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONTE = "Arial"

PALETA = ["F5C800", "3D3D3D", "8C8C8C", "C99A00", "BDBDBD",
          "5E5E5E", "FFE066", "A67C00", "E0E0E0", "1F1F1F"]

LARGURA, ALTURA = Inches(10), Inches(5.625)


def _rgb(hexa):
    return RGBColor.from_string(hexa.lstrip("#").upper())


# --------------------------------------------------------------------------
# Primitivas de desenho
# --------------------------------------------------------------------------

def _fundo(slide, cor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = cor


def _retangulo(slide, x, y, w, h, cor):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = cor
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _texto(slide, x, y, w, h, texto, tamanho=12, cor=GRAFITE_MEDIO,
           negrito=False, alinhamento=PP_ALIGN.LEFT, ancora=MSO_ANCHOR.TOP):
    caixa = slide.shapes.add_textbox(x, y, w, h)
    tf = caixa.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancora
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, linha in enumerate(str(texto).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinhamento
        run = p.add_run()
        run.text = linha
        f = run.font
        f.name, f.size, f.bold = FONTE, Pt(tamanho), negrito
        f.color.rgb = cor
    return caixa


def _cabecalho(slide, titulo, subtitulo=None):
    """Faixa grafite no topo com o título em amarelo (slides de conteúdo)."""
    _fundo(slide, BRANCO)
    _retangulo(slide, 0, 0, LARGURA, Inches(0.85), GRAFITE_MEDIO)
    _texto(slide, Inches(0.4), Inches(0.12), Inches(9.2), Inches(0.6),
           titulo.upper(), 24 if len(titulo) <= 38 else 20, AMARELO, True,
           ancora=MSO_ANCHOR.MIDDLE)
    if subtitulo:
        _texto(slide, Inches(0.4), Inches(5.2), Inches(9.2), Inches(0.3),
               subtitulo, 9, CINZA_TEXTO)


def _logo(slide, logo_path, x, y, altura):
    if logo_path and os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, x, y, height=altura)


def _card_kpi(slide, x, y, valor, rotulo, detalhe=None):
    """Card escuro com número grande em amarelo (quadro resumo)."""
    w, h = Inches(2.18), Inches(1.55)
    _retangulo(slide, x, y, w, h, CARD_ESCURO)
    _retangulo(slide, x, y, w, Inches(0.07), AMARELO)
    _texto(slide, x + Inches(0.1), y + Inches(0.18), w - Inches(0.2), Inches(0.62),
           valor, 26, AMARELO, True)
    _texto(slide, x + Inches(0.1), y + Inches(0.82), w - Inches(0.2), Inches(0.35),
           rotulo, 11, CINZA_SUAVE)
    if detalhe:
        _texto(slide, x + Inches(0.1), y + Inches(1.12), w - Inches(0.2), Inches(0.35),
               detalhe, 9, CINZA_SUAVE)


def _card_categoria(slide, x, y, w, h, titulo, detalhe, cor_barra):
    """Card claro com barra de cor lateral (legenda das distribuições)."""
    _retangulo(slide, x, y, w, h, CINZA_CLARO)
    _retangulo(slide, x, y, Inches(0.12), h, cor_barra)
    _texto(slide, x + Inches(0.25), y + Inches(0.05), w - Inches(0.35), h * 0.45,
           titulo, 13, GRAFITE_MEDIO, True, ancora=MSO_ANCHOR.BOTTOM)
    _texto(slide, x + Inches(0.25), y + h * 0.5, w - Inches(0.35), h * 0.45,
           detalhe, 11, CINZA_TEXTO)


def _tabela(slide, x, y, w, linhas, larguras=None, tamanho=11, alt_linha=0.32,
            escuro=False, colunas_texto=1):
    """Tabela com cabeçalho amarelo. `linhas[0]` é o cabeçalho.
    As `colunas_texto` primeiras colunas ficam à esquerda; as demais (números) à direita."""
    n_lin, n_col = len(linhas), len(linhas[0])
    shp = slide.shapes.add_table(n_lin, n_col, x, y, w, Inches(alt_linha * n_lin))
    tabela = shp.table
    if larguras:
        total = sum(larguras)
        for i, lw in enumerate(larguras):
            tabela.columns[i].width = int(w * lw / total)
    for r, linha in enumerate(linhas):
        tabela.rows[r].height = Inches(alt_linha)
        for c, valor in enumerate(linha):
            cel = tabela.cell(r, c)
            cel.fill.solid()
            if r == 0:
                cel.fill.fore_color.rgb = AMARELO
                cor_txt = GRAFITE
            elif escuro:
                cel.fill.fore_color.rgb = CARD_ESCURO if r % 2 else GRAFITE_MEDIO
                cor_txt = BRANCO
            else:
                cel.fill.fore_color.rgb = CINZA_CLARO if r % 2 else BRANCO
                cor_txt = GRAFITE_MEDIO
            cel.margin_left = cel.margin_right = Inches(0.08)
            cel.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cel.text_frame
            tf.text = str(valor)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if c < colunas_texto else PP_ALIGN.RIGHT
            for run in p.runs:
                run.font.name, run.font.size = FONTE, Pt(tamanho)
                run.font.bold = r == 0
                run.font.color.rgb = cor_txt
    return tabela


# --------------------------------------------------------------------------
# Gráficos nativos
# --------------------------------------------------------------------------

def _estilo_fonte_grafico(chart, tamanho=10):
    chart.has_title = False
    chart.font.name = FONTE
    chart.font.size = Pt(tamanho)
    chart.font.color.rgb = GRAFITE_MEDIO


def _grafico_rosca(slide, x, y, w, h, categorias, valores, cores):
    dados = CategoryChartData()
    dados.categories = categorias
    dados.add_series("% da extensão", valores)
    chart = slide.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, x, y, w, h, dados).chart
    _estilo_fonte_grafico(chart)
    chart.has_legend = False
    serie = chart.plots[0].series[0]
    for i, cor in enumerate(cores):
        ponto = serie.points[i]
        ponto.format.fill.solid()
        ponto.format.fill.fore_color.rgb = _rgb(cor)
        ponto.format.line.color.rgb = BRANCO
    plot = chart.plots[0]
    plot.has_data_labels = True
    rot = plot.data_labels
    rot.number_format = '0.0"%"'
    rot.number_format_is_linked = False
    rot.show_value = True
    rot.font.size = Pt(10)
    rot.font.bold = True
    rot.font.color.rgb = BRANCO
    return chart


def _grafico_barras(slide, x, y, w, h, categorias, valores, rotulos,
                    horizontal=True, cor="F5C800", titulo_serie="% da extensão"):
    """Barras com rótulo personalizado em cada barra (ex.: '53,6% (60,99 km)')."""
    dados = CategoryChartData()
    dados.categories = categorias
    dados.add_series(titulo_serie, valores)
    tipo = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    chart = slide.shapes.add_chart(tipo, x, y, w, h, dados).chart
    _estilo_fonte_grafico(chart, 9)
    chart.has_legend = False

    plot = chart.plots[0]
    plot.gap_width = 45
    serie = plot.series[0]
    serie.format.fill.solid()
    serie.format.fill.fore_color.rgb = _rgb(cor)

    eixo_cat = chart.category_axis
    eixo_cat.format.line.color.rgb = CINZA_SUAVE
    eixo_cat.has_major_gridlines = False
    eixo_cat.tick_labels.font.size = Pt(9)
    if horizontal:
        eixo_cat.reverse_order = True  # maior no topo

    eixo_val = chart.value_axis
    eixo_val.visible = False
    eixo_val.has_major_gridlines = False
    eixo_val.maximum_scale = max(valores) * 1.35 if valores else 1
    eixo_val.minimum_scale = 0

    for i, texto in enumerate(rotulos):
        rot = serie.points[i].data_label
        rot.has_text_frame = True
        rot.text_frame.text = texto
        rot.position = XL_LABEL_POSITION.OUTSIDE_END
        for par in rot.text_frame.paragraphs:
            for run in par.runs:
                run.font.size = Pt(9)
                run.font.name = FONTE
                run.font.color.rgb = GRAFITE_MEDIO
    return chart


# --------------------------------------------------------------------------
# Slides
# --------------------------------------------------------------------------

def _slide_vazio(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])  # layout em branco


def slide_capa(prs, rel, logo_path):
    s = _slide_vazio(prs)
    _fundo(s, GRAFITE)
    _retangulo(s, 0, 0, Inches(0.18), ALTURA, AMARELO)
    if logo_path and os.path.exists(logo_path):
        _retangulo(s, Inches(0.6), Inches(0.5), Inches(2.6), Inches(0.9), BRANCO)
        _logo(s, logo_path, Inches(0.75), Inches(0.58), Inches(0.74))
    _texto(s, Inches(0.6), Inches(1.9), Inches(8.8), Inches(0.5),
           "RELATÓRIO TÉCNICO", 14, CINZA_SUAVE, True)
    _texto(s, Inches(0.6), Inches(2.35), Inches(8.8), Inches(0.9),
           "MALHA VIÁRIA MUNICIPAL", 40, AMARELO, True)
    _retangulo(s, Inches(0.6), Inches(3.35), Inches(1.2), Inches(0.06), AMARELO)
    _texto(s, Inches(0.6), Inches(3.55), Inches(8.8), Inches(0.5),
           f"Relatório de Logradouros · {rel['municipio']}", 18, BRANCO)
    _texto(s, Inches(0.6), Inches(4.05), Inches(8.8), Inches(0.4),
           rel["data"], 13, CINZA_SUAVE)


def slide_resumo(prs, rel, df, col_map):
    s = _slide_vazio(prs)
    _fundo(s, GRAFITE)
    _retangulo(s, 0, 0, Inches(0.18), ALTURA, AMARELO)
    _texto(s, Inches(0.35), Inches(0.22), Inches(9.3), Inches(0.65),
           "MALHA VIÁRIA MUNICIPAL", 30, AMARELO, True)
    _texto(s, Inches(0.35), Inches(0.85), Inches(9.3), Inches(0.35),
           f"Relatório de Logradouros · {rel['municipio']} · {rel['data']}", 13, CINZA_SUAVE)

    kpis = [(formatar_inteiro_br(rel.get("total_trechos", rel["total_registros"])),
             "Trechos mapeados", None)]
    if rel.get("total_vias"):
        kpis.append((formatar_inteiro_br(rel["total_vias"]), "Logradouros", "vias distintas"))
    if "total_km" in rel:
        kpis.append((formatar_numero_br(rel["total_km"], 2), "km de extensão", "malha total"))
    if "pct_sem_nome_vias" in rel:
        kpis.append((formatar_pct(rel["pct_sem_nome_vias"], 1), "Sem denominação",
                     f"{formatar_pct(rel.get('pct_sem_nome_ext', 0), 1)} da extensão"))
    for i, (valor, rotulo, det) in enumerate(kpis[:4]):
        _card_kpi(s, Inches(0.35 + i * 2.38), Inches(1.35), valor, rotulo, det)

    linhas = [["INDICADOR", "VALOR"]]
    comp = col_map.get("comprimento")
    if comp and comp in df.columns:
        serie = df.loc[df[comp] > 0, comp]
        if not serie.empty:
            linhas += [
                ["Comprimento mínimo de trecho (m)", formatar_numero_br(serie.min())],
                ["Comprimento máximo de trecho (m)", formatar_numero_br(serie.max())],
                ["Mediana do comprimento (m)", formatar_numero_br(serie.median())],
            ]
    if col_map.get("bairro"):
        n = filtrar_bairros_validos(df, col_map)[col_map["bairro"]].nunique()
        linhas.append(["Total de bairros mapeados", f"{formatar_inteiro_br(n)} bairros"])
    if col_map.get("setor"):
        n = df[col_map["setor"]].nunique()
        linhas.append(["Total de setores", f"{formatar_inteiro_br(n)} setores"])
    if len(linhas) > 1:
        _tabela(s, Inches(0.35), Inches(3.1), Inches(9.3), linhas[:6],
                larguras=[3, 1], tamanho=11, alt_linha=0.33, escuro=True)


def _cores_categorias(indices, por_tipo):
    if por_tipo:
        return [get_cor_pavimentacao(c).lstrip("#") for c in indices]
    return [PALETA[i % len(PALETA)] for i in range(len(indices))]


def slide_distribuicao_rosca(prs, tabela, titulo, por_tipo=True, max_cards=4):
    """Rosca à esquerda + cards de categoria à direita (padrão do slide de pavimentação)."""
    s = _slide_vazio(prs)
    _cabecalho(s, titulo, "Participação na extensão da malha · km da categoria ÷ km total")

    cats = [str(c) for c in tabela.index]
    pct = [float(v) for v in tabela["% da Extensão"]]
    cores = _cores_categorias(cats, por_tipo)
    _grafico_rosca(s, Inches(0.3), Inches(1.0), Inches(4.6), Inches(4.1), cats, pct, cores)

    n = min(len(cats), max_cards)
    area_h = 3.9
    gap = 0.15
    card_h = min(1.1, (area_h - gap * (n - 1)) / max(n, 1))
    for i in range(n):
        km = tabela["_extensao_m"].iloc[i] / 1000
        trechos = tabela["_trechos"].iloc[i]
        det = (f"{formatar_numero_br(km)} km  ·  {formatar_inteiro_br(trechos)} trechos"
               f"  ·  {formatar_pct(pct[i])}")
        _card_categoria(s, Inches(5.3), Inches(1.1 + i * (card_h + gap)), Inches(4.3),
                        Inches(card_h), cats[i].title(), det, _rgb(cores[i]))
    if len(cats) > max_cards:
        resto = sum(pct[max_cards:])
        _texto(s, Inches(5.3), Inches(1.1 + n * (card_h + gap)), Inches(4.3), Inches(0.3),
               f"+ {len(cats) - max_cards} outras categorias ({formatar_pct(resto)})",
               10, CINZA_TEXTO)


def slide_distribuicao_barras(prs, tabela, titulo, top=10, legenda=None):
    s = _slide_vazio(prs)
    tab = tabela.head(top)
    sub = legenda or "Participação na extensão da malha · km ÷ km total"
    if len(tabela) > top:
        sub += f" · exibindo {top} de {len(tabela)}"
    _cabecalho(s, titulo, sub)
    cats = [str(c).title() for c in tab.index]
    pct = [float(v) for v in tab["% da Extensão"]]
    rot = [f"{formatar_pct(p, 1)} ({formatar_numero_br(m / 1000)} km)"
           for p, m in zip(pct, tab["_extensao_m"])]
    _grafico_barras(s, Inches(0.4), Inches(1.0), Inches(9.2), Inches(4.15),
                    cats, pct, rot, horizontal=len(cats) > 5)


def slide_denominacao(prs, rel):
    geral = rel["denominacao_geral"]
    s = _slide_vazio(prs)
    _cabecalho(s, "Denominação dos logradouros",
               "% das vias distintas · sem denominação: 'RUA PROJETADA', 'SEM NOME', nome em branco "
               "ou só tipo + número (ex.: TRAVESSA 2)")

    com, sem = geral.loc["Com denominação"], geral.loc["Sem denominação"]
    if int(sem["Vias (qtd)"]) == 0:
        # Rosca de 100% x 0% não diz nada: mostra o resultado em destaque
        _retangulo(s, Inches(0.4), Inches(1.3), Inches(9.2), Inches(3.4), CINZA_CLARO)
        _retangulo(s, Inches(0.4), Inches(1.3), Inches(0.12), Inches(3.4), AMARELO)
        _texto(s, Inches(0.8), Inches(1.6), Inches(8.5), Inches(1.1),
               "100%", 60, GRAFITE_MEDIO, True)
        _texto(s, Inches(0.8), Inches(2.75), Inches(8.5), Inches(0.5),
               "dos logradouros possuem denominação", 20, GRAFITE_MEDIO, True)
        _texto(s, Inches(0.8), Inches(3.35), Inches(8.5), Inches(0.9),
               f"{formatar_inteiro_br(com['Vias (qtd)'])} logradouros  ·  "
               f"{formatar_numero_br(com['Extensão (km)'])} km. Nenhuma via registrada "
               "como projetada, 'SEM NOME' ou com o nome em branco.", 13, CINZA_TEXTO)
        return
    _grafico_rosca(s, Inches(0.3), Inches(1.0), Inches(4.6), Inches(4.1),
                   ["Com denominação", "Sem denominação"],
                   [float(com["% das Vias"]), float(sem["% das Vias"])],
                   ["3D3D3D", "F5C800"])

    for i, (linha, nome, cor) in enumerate([(com, "Com denominação", GRAFITE_MEDIO),
                                            (sem, "Sem denominação", AMARELO)]):
        det = (f"{formatar_inteiro_br(linha['Vias (qtd)'])} logradouros  ·  "
               f"{formatar_pct(linha['% das Vias'])}\n"
               f"{formatar_numero_br(linha['Extensão (km)'])} km  ·  "
               f"{formatar_pct(linha['% da Extensão'])} da extensão")
        _card_categoria(s, Inches(5.3), Inches(1.1 + i * 1.45), Inches(4.3), Inches(1.3),
                        nome, det, cor)
    _texto(s, Inches(5.3), Inches(4.1), Inches(4.3), Inches(0.9),
           "Ter nome é uma propriedade do logradouro inteiro, por isso a base é a "
           "quantidade de vias; a extensão aparece como referência.", 10, CINZA_TEXTO)


def slide_bairros_sem_nome(prs, rank, top=10):
    s = _slide_vazio(prs)
    tab = rank.head(top)
    sub = "Quantidade de logradouros (ruas distintas) sem nome por bairro"
    if len(rank) > top:
        sub += f" · exibindo {top} de {len(rank)}"
    _cabecalho(s, "Bairros com mais logradouros sem denominação", sub)
    cats = [str(c).title() for c in tab.index]
    vals = [int(v) for v in tab["Logradouros sem denominação"]]
    _grafico_barras(s, Inches(0.4), Inches(1.0), Inches(9.2), Inches(4.15), cats, vals,
                    [formatar_inteiro_br(v) for v in vals], horizontal=True,
                    cor="3D3D3D", titulo_serie="Logradouros sem denominação")


def slide_top_vias(prs, top_vias):
    s = _slide_vazio(prs)
    _cabecalho(s, "As 10 vias mais extensas", "Extensão somada de todos os trechos de cada via")
    linhas = [["#", "VIA", "EXTENSÃO (km)", "% DA MALHA"]]
    for pos, r in top_vias.iterrows():
        linhas.append([str(pos), str(r["Via"]).title(), formatar_numero_br(r["Extensão (km)"], 3),
                       formatar_pct(r["% da Extensão"])])
    _tabela(s, Inches(0.4), Inches(1.05), Inches(9.2), linhas,
            larguras=[0.5, 5, 1.6, 1.4], tamanho=10, alt_linha=0.36, colunas_texto=2)


def slide_encerramento(prs, rel, logo_path):
    s = _slide_vazio(prs)
    _fundo(s, GRAFITE)
    _retangulo(s, 0, 0, Inches(0.18), ALTURA, AMARELO)
    _texto(s, Inches(0.6), Inches(1.9), Inches(8.8), Inches(0.9),
           "OBRIGADO", 44, AMARELO, True, PP_ALIGN.CENTER)
    _texto(s, Inches(0.6), Inches(2.9), Inches(8.8), Inches(0.5),
           f"Malha Viária Municipal · {rel['municipio']} · {rel['data']}", 14, CINZA_SUAVE,
           alinhamento=PP_ALIGN.CENTER)
    if logo_path and os.path.exists(logo_path):
        _retangulo(s, Inches(3.7), Inches(3.8), Inches(2.6), Inches(0.9), BRANCO)
        _logo(s, logo_path, Inches(3.85), Inches(3.88), Inches(0.74))


# --------------------------------------------------------------------------
# Entrada principal
# --------------------------------------------------------------------------

def gerar_apresentacao(df, col_map, rel, caminho, logo_path=None):
    prs = Presentation()
    prs.slide_width, prs.slide_height = LARGURA, ALTURA

    slide_capa(prs, rel, logo_path)
    slide_resumo(prs, rel, df, col_map)
    if rel.get("por_pavimentacao") is not None:
        slide_distribuicao_rosca(prs, rel["por_pavimentacao"], "Distribuição por pavimentação")
    if rel.get("por_status") is not None:
        slide_distribuicao_rosca(prs, rel["por_status"], "Status das vias")
    if rel.get("por_setor") is not None:
        slide_distribuicao_barras(prs, rel["por_setor"], "Distribuição por setor", top=12)
    if rel.get("por_bairro") is not None:
        slide_distribuicao_barras(prs, rel["por_bairro"], "Bairros com maior extensão", top=10)
    if rel.get("denominacao_geral") is not None:
        slide_denominacao(prs, rel)
    if rel.get("ranking_bairro_sem_nome") is not None:
        slide_bairros_sem_nome(prs, rel["ranking_bairro_sem_nome"])
    if rel.get("top_vias") is not None:
        slide_top_vias(prs, rel["top_vias"])
    slide_encerramento(prs, rel, logo_path)

    prs.save(caminho)
    return caminho
